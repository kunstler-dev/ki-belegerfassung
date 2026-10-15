"""Export of the results as an Excel workbook and as a CSV file for import into accounting software."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Ergebnis

SPALTEN = ["Datei", "Status", "Lieferant", "USt-IdNr.", "Rechnungsnr.", "Rechnungsdatum", "Fällig am",
           "Netto", "MwSt. %", "MwSt.", "Brutto", "Währung", "IBAN", "Hinweise"]
FARBEN = {"OK": "D9EAD3", "Prüfen": "FFF2CC", "Fehler": "F4CCCC"}


def zeile(e: Ergebnis) -> list:
    b = e.beleg
    if b is None:
        return [e.datei, e.status] + [None] * 11 + [e.fehler]
    return [e.datei, e.status, b.lieferant, b.lieferant_ust_id, b.rechnungsnummer, b.rechnungsdatum,
            b.faelligkeitsdatum, float(b.netto), float(b.mwst_satz) if b.mwst_satz is not None else None,
            float(b.mwst_betrag), float(b.brutto), b.waehrung, b.iban, " ".join(e.hinweise) or None]


def als_excel(ergebnisse: list[Ergebnis], ziel: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Belege"
    ws.append(SPALTEN)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
    for e in ergebnisse:
        ws.append(zeile(e))
        ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=FARBEN[e.status])
    for spalte in ("F", "G"):
        for zelle in ws[spalte][1:]:
            zelle.number_format = "DD.MM.YYYY"
    for spalte in ("H", "J", "K"):
        for zelle in ws[spalte][1:]:
            zelle.number_format = "#,##0.00"
    for i, name in enumerate(SPALTEN, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(name) + 2)
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["N"].width = 60
    ws.freeze_panes = "A2"
    wb.save(ziel)
    return ziel


def als_csv(ergebnisse: list[Ergebnis], ziel: Path) -> Path:
    """Semicolon-separated with decimal commas: what German Excel and most accounting tools expect."""
    with ziel.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(SPALTEN)
        for e in ergebnisse:
            w.writerow([_de(v) for v in zeile(e)])
    return ziel


def _de(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float):
        return f"{wert:.2f}".replace(".", ",")
    if hasattr(wert, "strftime"):
        return wert.strftime("%d.%m.%Y")
    return str(wert)
