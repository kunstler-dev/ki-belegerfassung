import base64
import csv
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from types import SimpleNamespace

import pytest
from anthropic.lib._parse._transform import transform_schema
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from belegerfassung import cli, extract
from belegerfassung.checks import iban_gueltig, pruefen
from belegerfassung.export import als_csv, als_excel
from belegerfassung.models import Beleg, Ergebnis, Position
from belegerfassung.web import create_app

BEISPIELE = Path(__file__).resolve().parent.parent / "beispiele"
HEUTE = date(2026, 10, 15)


def beleg(**aenderungen) -> Beleg:
    daten = dict(
        lieferant="Papierhaus Muster GmbH", lieferant_ust_id="DE123456789", rechnungsnummer="PM-1",
        rechnungsdatum=date(2026, 10, 2), faelligkeitsdatum=date(2026, 10, 16), waehrung="EUR",
        netto=D("220.00"), mwst_satz=D("19"), mwst_betrag=D("41.80"), brutto=D("261.80"),
        iban="DE89 3704 0044 0532 0130 00",
        positionen=[Position(beschreibung="Papier", menge=D("10"), einzelpreis=D("4.90"), gesamt=D("49.00")),
                    Position(beschreibung="Toner", menge=D("2"), einzelpreis=D("85.50"), gesamt=D("171.00"))],
    )
    daten.update(aenderungen)
    return Beleg(**daten)


# ---------------- checks ----------------

def test_consistent_invoice_has_no_findings():
    assert pruefen(beleg(), HEUTE) == []


def test_gross_mismatch_is_flagged():
    hinweise = pruefen(beleg(brutto=D("271.80")), HEUTE)
    assert any("Bruttobetrag" in h for h in hinweise)


def test_vat_amount_must_match_rate():
    hinweise = pruefen(beleg(mwst_betrag=D("15.40"), brutto=D("235.40")), HEUTE)
    assert any("passt nicht zu 19" in h for h in hinweise)


def test_line_items_must_add_up_to_net():
    hinweise = pruefen(beleg(positionen=[Position(beschreibung="x", menge=None, einzelpreis=None, gesamt=D("100"))]), HEUTE)
    assert any("Positionen" in h for h in hinweise)


def test_rounding_within_two_cents_is_accepted():
    assert pruefen(beleg(brutto=D("261.81")), HEUTE) == []


def test_dates_and_currency():
    hinweise = pruefen(beleg(rechnungsdatum=date(2026, 12, 1), faelligkeitsdatum=date(2026, 11, 1), waehrung="USD"), HEUTE)
    assert len(hinweise) == 3


def test_mixed_vat_rates_skip_rate_check():
    assert pruefen(beleg(mwst_satz=None), HEUTE) == []


@pytest.mark.parametrize("iban,ok", [
    ("DE89 3704 0044 0532 0130 00", True),
    ("DE02120300000000202051", True),
    ("DE89 3704 0044 0532 0130 01", False),
    ("kein iban", False),
])
def test_iban_checksum(iban, ok):
    assert iban_gueltig(iban) is ok


# ---------------- API request ----------------

class FakeMessages:
    def __init__(self, antwort):
        self.antwort = antwort
        self.calls = []

    def parse(self, **kwargs):
        self.calls.append(kwargs)
        return self.antwort


def fake_client(parsed=None, stop_reason="end_turn"):
    messages = FakeMessages(SimpleNamespace(parsed_output=parsed, stop_reason=stop_reason))
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


def test_pdf_is_sent_as_document_with_schema_and_fallback():
    client, messages = fake_client(beleg())
    pfad = BEISPIELE / "rechnung-papierhaus.pdf"
    assert extract.lesen(client, pfad) == beleg()

    call = messages.calls[0]
    assert call["model"] == "claude-opus-5-5"
    assert call["output_format"] is Beleg
    assert call["fallbacks"] == "default" and call["betas"] == ["server-side-fallback-2026-07-01"]
    block = call["messages"][0]["content"][0]
    assert block["type"] == "document" and block["source"]["media_type"] == "application/pdf"
    assert base64.standard_b64decode(block["source"]["data"]) == pfad.read_bytes()


def test_scan_is_sent_as_image():
    client, messages = fake_client(beleg())
    extract.lesen(client, BEISPIELE / "rechnung-druckerei-scan.png")
    block = messages.calls[0]["messages"][0]["content"][0]
    assert block["type"] == "image" and block["source"]["media_type"] == "image/png"


def test_refusal_and_unsupported_files_become_errors(tmp_path):
    client, _ = fake_client(None, stop_reason="refusal")
    e = extract.verarbeiten(client, BEISPIELE / "rechnung-papierhaus.pdf")
    assert e.status == "Fehler" and "abgelehnt" in e.fehler

    txt = tmp_path / "notiz.txt"
    txt.write_text("hallo")
    e = extract.verarbeiten(client, txt)
    assert e.status == "Fehler" and "Dateiformat" in e.fehler


def test_schema_is_accepted_by_structured_outputs():
    """The SDK converts the Pydantic model into the JSON schema subset the API supports."""
    schema = transform_schema(Beleg.model_json_schema())
    assert schema["additionalProperties"] is False
    assert set(schema["required"]) == set(Beleg.model_fields)


# ---------------- export ----------------

def demo() -> list[Ergebnis]:
    return cli.demo_ergebnisse()


def test_demo_results_match_the_checks():
    """The bundled demo data must be what the checks really produce for those invoices."""
    for e in demo():
        assert e.hinweise == pruefen(e.beleg, HEUTE)
    status = {e.datei: e.status for e in demo()}
    assert status["rechnung-it-service.pdf"] == "Prüfen"
    assert status["rechnung-papierhaus.pdf"] == "OK"


def test_excel_export(tmp_path):
    ziel = als_excel(demo() + [Ergebnis(datei="kaputt.pdf", fehler="API-Fehler")], tmp_path / "b.xlsx")
    ws = load_workbook(ziel).active
    assert ws["A1"].value == "Datei" and ws.max_row == 5
    werte = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=11).value for r in range(2, 6)}
    assert werte["rechnung-papierhaus.pdf"] == 261.8


def test_csv_uses_german_number_format(tmp_path):
    ziel = als_csv(demo(), tmp_path / "b.csv")
    rows = list(csv.reader(ziel.open(encoding="utf-8-sig"), delimiter=";"))
    papier = next(r for r in rows if r[0] == "rechnung-papierhaus.pdf")
    assert papier[10] == "261,80" and papier[5] == "02.10.2026"


# ---------------- CLI and web ----------------

def test_cli_demo_writes_files(tmp_path, capsys):
    assert cli.main(["einlesen", "--demo", "-o", str(tmp_path / "out")]) == 0
    assert (tmp_path / "out.xlsx").exists() and (tmp_path / "out.csv").exists()
    assert "3 Belege: 2 OK, 1 zu prüfen" in capsys.readouterr().out


def test_web_upload_and_export():
    client, _ = fake_client(beleg())
    app = TestClient(create_app(client=client))
    assert "KI-<span>Belegerfassung" in app.get("/").text

    pdf = (BEISPIELE / "rechnung-papierhaus.pdf").read_bytes()
    r = app.post("/api/belege", files=[("dateien", ("r.pdf", pdf, "application/pdf")),
                                       ("dateien", ("bild.bmp", b"x", "image/bmp"))])
    daten = r.json()
    assert [d["status"] for d in daten] == ["OK", "Fehler"]
    assert app.get("/export/xlsx").status_code == 200
    assert app.get("/export/pdf").status_code == 404


def test_web_demo_mode():
    app = TestClient(create_app(demo=True))
    assert len(app.get("/api/belege").json()) == 3
    assert app.post("/api/belege", files=[("dateien", ("r.pdf", b"%PDF", "application/pdf"))]).status_code == 400
    assert json.loads(app.get("/api/belege").content)[0]["beleg"]["lieferant"]
